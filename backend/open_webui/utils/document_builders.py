"""Builds pdf, docx, xlsx and pptx files, and runs as a child process (python -m) so a hung build can be killed."""

import json
import logging
import sys
import threading
from functools import lru_cache
from pathlib import Path

log = logging.getLogger(__name__)

# Exit code that tells the parent the file came out over its size limit
EXIT_TOO_LARGE = 3

# Set by main() from what the parent process passes in, so the child doesn't import the whole app config for it
_paths: dict[str, str] = {}

# Set by main() when the parent has no restyled reference.docx, so the child doesn't try to build one
_skip_reference_doc = False


def _fonts_dir_setting():
    if 'fonts_dir' in _paths:
        return Path(_paths['fonts_dir'])
    from open_webui.env import FONTS_DIR

    return FONTS_DIR


@lru_cache(maxsize=None)
def _resolve_static_asset_dir(configured_path, relative_subpath: str):
    """Fall back from the configured path to site-packages, then a repo-relative path (installed vs run from backend/)."""
    import site
    from pathlib import Path

    for candidate in (
        configured_path,
        Path(site.getsitepackages()[0]) / relative_subpath,
        Path('.') / 'backend' / relative_subpath,
    ):
        if candidate.exists():
            return candidate
    return configured_path


PDF_TEXT_COLOR = '#000000'
# Google Docs' 1in page margins, in fpdf's millimeters
PDF_PAGE_MARGIN = 25.4
MM_PER_PT = 25.4 / 72
PDF_BLOCK_TAG_WHITESPACE = (
    r'\s*(</?(?:p|h[1-6]|ul|ol|li|blockquote|pre|table|thead|tbody|tr|th|td|hr|br|div)\b[^>]*>)\s*'
)


def _build_pdf_document_bytes(title: str, content: str) -> bytes:
    import re
    from html import escape

    from fpdf import FPDF, FontFace, TextStyle
    from fpdf.html import HTML2FPDF
    from markdown import markdown

    class DocumentHTML2FPDF(HTML2FPDF):
        def handle_starttag(self, tag, attrs):
            # fpdf draws list bullets in the page's last font, so a list under a heading got heading-sized numbers
            if tag == 'li':
                self.pdf.set_font(family=self.font_family, size=self.font_size_pt, style=self.font_emphasis.style)
            super().handle_starttag(tag, attrs)

    class DocumentPDF(FPDF):
        HTML2FPDF_CLASS = DocumentHTML2FPDF

    fonts_dir = _resolve_static_asset_dir(_fonts_dir_setting(), 'static/fonts')

    pdf = DocumentPDF()
    pdf.set_margins(PDF_PAGE_MARGIN, PDF_PAGE_MARGIN, PDF_PAGE_MARGIN)
    pdf.add_page()
    pdf.add_font('NotoSans', '', f'{fonts_dir}/NotoSans-Regular.ttf')
    pdf.add_font('NotoSans', 'b', f'{fonts_dir}/NotoSans-Bold.ttf')
    pdf.add_font('NotoSans', 'i', f'{fonts_dir}/NotoSans-Italic.ttf')
    # There's no bold italic file, and without this ***text*** fails the whole build
    pdf.add_font('NotoSans', 'bi', f'{fonts_dir}/NotoSans-Bold.ttf')
    pdf.set_font('NotoSans', size=11)
    pdf.set_text_color(PDF_TEXT_COLOR)
    pdf.set_auto_page_break(auto=True, margin=PDF_PAGE_MARGIN)

    # The default Courier for code has no diacritics, so "ključ" in backticks would fail the build
    code_style = FontFace(family='NotoSans')
    tag_styles = {
        'code': code_style,
        'pre': code_style,
        # fpdf ignores a list item's top margin, so items are spaced with a bottom one
        'li': TextStyle(t_margin=2, l_margin=5, b_margin=1.5),
        # With no top margin of their own, lists get fpdf's paragraph gap, and none right under a heading
        'ul': TextStyle(t_margin=0),
        'ol': TextStyle(t_margin=0),
        'blockquote': TextStyle(color='#374151', t_margin=3, b_margin=3, l_margin=6),
        # fpdf's defaults are dark red headings, these are the same Google Docs headings our docx uses
        **{
            f'h{level}': TextStyle(
                font_style='I' if f'Heading{level}' in REFERENCE_DOC_ITALIC_HEADINGS else '',
                font_size_pt=size,
                color=f'#{color}',
                t_margin=before * MM_PER_PT,
                b_margin=after * MM_PER_PT,
            )
            for level in range(1, 7)
            for size, color, before, after in [REFERENCE_DOC_HEADINGS[f'Heading{level}']]
        },
    }
    # Models usually open with their own "# Title", so ours is only added when the content has none
    heading = '' if content.lstrip().startswith('#') else f'<h1>{escape(title)}</h1>'
    # fpdf renders the newlines markdown leaves around block tags as blank lines, spaces between inline tags must stay
    html = re.sub(PDF_BLOCK_TAG_WHITESPACE, r'\1', heading + markdown(content))
    pdf.write_html(html, tag_styles=tag_styles, li_prefix_color=PDF_TEXT_COLOR)
    return bytes(pdf.output())


REFERENCE_DOC_TABLE_BORDER = '<w:{side} w:val="single" w:sz="4" w:space="0" w:color="BFBFBF" />'

# Google Docs' default headings as (size in pt, color, space before and after in pt), none of them bold
REFERENCE_DOC_HEADINGS = {
    'Title': (26, '000000', 0, 3),
    'Subtitle': (15, '666666', 0, 16),
    'Heading1': (20, '000000', 20, 6),
    'Heading2': (16, '000000', 18, 6),
    'Heading3': (14, '434343', 16, 4),
    'Heading4': (12, '666666', 14, 4),
    'Heading5': (11, '666666', 12, 4),
    'Heading6': (11, '666666', 12, 4),
}
REFERENCE_DOC_ITALIC_HEADINGS = {'Heading6'}


def _reference_doc_heading_style(style_id: str) -> str:
    size, color, before, after = REFERENCE_DOC_HEADINGS[style_id]
    level = style_id.removeprefix('Heading')
    name = f'heading {level}' if level.isdigit() else style_id
    outline = f'<w:outlineLvl w:val="{int(level) - 1}" />' if level.isdigit() else ''
    return (
        f'<w:style w:type="paragraph" w:styleId="{style_id}"><w:name w:val="{name}" />'
        f'<w:basedOn w:val="Normal" /><w:next w:val="BodyText" /><w:link w:val="{style_id}Char" />'
        '<w:uiPriority w:val="9" /><w:qFormat />'
        f'<w:pPr><w:keepNext /><w:keepLines /><w:spacing w:before="{before * 20}" w:after="{after * 20}" />'
        f'{outline}</w:pPr>'
        f'<w:rPr>{"<w:i />" if style_id in REFERENCE_DOC_ITALIC_HEADINGS else ""}'
        f'<w:color w:val="{color}" /><w:sz w:val="{size * 2}" /><w:szCs w:val="{size * 2}" /></w:rPr>'
        '</w:style>'
    )


# Blocks swapped by pattern, since re-serialising with ElementTree would rename the namespace prefixes
REFERENCE_DOC_REPLACEMENTS = {
    'word/styles.xml': [
        (
            r'<w:docDefaults>.*?</w:docDefaults>',
            '<w:docDefaults><w:rPrDefault><w:rPr>'
            '<w:rFonts w:ascii="Arial" w:eastAsia="Arial" w:hAnsi="Arial" w:cs="Arial" />'
            '<w:sz w:val="22" /><w:szCs w:val="22" />'
            '<w:lang w:val="en-US" w:eastAsia="zh-CN" w:bidi="ar-SA" />'
            '</w:rPr></w:rPrDefault>'
            '<w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="276" w:lineRule="auto" /></w:pPr></w:pPrDefault>'
            '</w:docDefaults>',
        ),
        (
            r'<w:style [^>]*w:styleId="BodyText".*?</w:style>',
            '<w:style w:type="paragraph" w:styleId="BodyText"><w:name w:val="Body Text" />'
            '<w:basedOn w:val="Normal" /><w:link w:val="BodyTextChar" /><w:qFormat />'
            '<w:pPr><w:spacing w:before="60" w:after="120" /></w:pPr></w:style>',
        ),
        (
            r'<w:style [^>]*w:styleId="Compact".*?</w:style>',
            '<w:style w:type="paragraph" w:customStyle="1" w:styleId="Compact"><w:name w:val="Compact" />'
            '<w:basedOn w:val="BodyText" /><w:qFormat />'
            '<w:pPr><w:spacing w:before="0" w:after="60" /></w:pPr></w:style>',
        ),
        (
            r'<w:style [^>]*w:styleId="Table".*?</w:style>',
            '<w:style w:type="table" w:default="1" w:styleId="Table"><w:name w:val="Table" />'
            '<w:basedOn w:val="TableNormal" /><w:qFormat /><w:tblPr><w:tblInd w:w="0" w:type="dxa" />'
            '<w:tblBorders>'
            + ''.join(
                REFERENCE_DOC_TABLE_BORDER.format(side=side)
                for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')
            )
            + '</w:tblBorders>'
            '<w:tblCellMar><w:top w:w="57" w:type="dxa" /><w:left w:w="108" w:type="dxa" />'
            '<w:bottom w:w="57" w:type="dxa" /><w:right w:w="108" w:type="dxa" /></w:tblCellMar></w:tblPr>'
            '<w:tblStylePr w:type="firstRow"><w:rPr><w:b /><w:bCs /></w:rPr>'
            '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="F2F2F2" /></w:tcPr></w:tblStylePr></w:style>',
        ),
        *(
            (rf'<w:style [^>]*w:styleId="{style_id}".*?</w:style>', _reference_doc_heading_style(style_id))
            for style_id in REFERENCE_DOC_HEADINGS
        ),
    ],
    # pandoc takes the page from here, and without a size Word falls back to US Letter
    'word/document.xml': [
        (
            r'<w:sectPr>.*?</w:sectPr>',
            '<w:sectPr><w:footnotePr><w:numRestart w:val="eachSect" /></w:footnotePr>'
            '<w:pgSz w:w="11906" w:h="16838" />'
            '<w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" '
            'w:header="720" w:footer="720" w:gutter="0" /></w:sectPr>',
        ),
    ],
    'word/theme/theme1.xml': [
        (r'(?<=<a:majorFont>)\s*<a:latin [^>]*/>', '<a:latin typeface="Arial"/>'),
        (r'(?<=<a:minorFont>)\s*<a:latin [^>]*/>', '<a:latin typeface="Arial"/>'),
    ],
}


def _apply_reference_doc_replacements(docx_bytes: bytes) -> bytes:
    """Restyle a reference.docx via REFERENCE_DOC_REPLACEMENTS, raising ValueError if a block is missing."""
    import io
    import re
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(docx_bytes))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in REFERENCE_DOC_REPLACEMENTS:
                text = data.decode('utf-8')
                for pattern, replacement in REFERENCE_DOC_REPLACEMENTS[item.filename]:
                    text, count = re.subn(pattern, lambda _: replacement, text, count=1, flags=re.S)
                    if not count:
                        raise ValueError(f'{pattern} not found in {item.filename}')
                data = text.encode('utf-8')
            out.writestr(item, data)
    return buf.getvalue()


# Docx files in one batch are built in parallel threads, so the cache is checked under the lock and they build it once
_reference_doc_lock = threading.Lock()
_reference_doc = None


def _get_pandoc_reference_doc():
    """The reference.docx for docx output - pandoc's default restyled once per process, or None if that fails."""
    global _reference_doc
    with _reference_doc_lock:
        # A failed build isn't kept, so a passing problem doesn't leave pandoc's defaults until a restart
        if _reference_doc is None:
            _reference_doc = _build_pandoc_reference_doc()
        return _reference_doc


def _build_pandoc_reference_doc():
    import os
    import subprocess

    import pypandoc

    from open_webui.config import CACHE_DIR

    try:
        default_doc = subprocess.run(
            [pypandoc.get_pandoc_path(), '--print-default-data-file', 'reference.docx'], capture_output=True, check=True
        ).stdout
        restyled = _apply_reference_doc_replacements(default_doc)

        path = CACHE_DIR / 'pandoc' / 'reference.docx'
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = path.with_suffix(f'.{os.getpid()}.tmp')
        tmp_path.write_bytes(restyled)
        os.replace(tmp_path, path)
        return path
    except Exception as e:
        log.warning(f'Could not build the docx reference styles, falling back to pandoc defaults: {e}')
        return None


def _build_docx_document_bytes(content: str) -> bytes:
    import os
    import tempfile

    import pypandoc

    reference_doc = None if _skip_reference_doc else _get_pandoc_reference_doc()
    extra_args = ['--reference-doc', str(reference_doc)] if reference_doc else []

    # The parent gives a directory it removes itself, since a killed build can't clean up after itself
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False, dir=_paths.get('tmp_dir')) as tmp:
        tmp_path = tmp.name
    try:
        # markdown-auto_identifiers avoids pandoc turning each heading into a stray Word bookmark
        pypandoc.convert_text(
            content, 'docx', format='markdown-auto_identifiers', outputfile=tmp_path, extra_args=extra_args
        )
        with open(tmp_path, 'rb') as f:
            return f.read()
    finally:
        os.unlink(tmp_path)


XLSX_FONT = ('Arial', 10)
# Google Sheets' default column width and row height, as its own xlsx export writes them
XLSX_MIN_COLUMN_WIDTH = 12.63
XLSX_ROW_HEIGHT = 15.75
XLSX_MAX_COLUMN_WIDTH = 60


def _xlsx_cell_value(text: str):
    """A CSV cell as a number with its format when it plainly is one, so "007" or "10.749,20" stay text."""
    import re

    value = text.strip()
    # Models copy Excel's "'" prefix for forcing text, which we'd otherwise write into the cell as a visible character
    if value.startswith("'"):
        return value[1:], None
    percent = value.endswith('%')
    number = value[:-1].strip() if percent else value
    if re.fullmatch(r'-?(0|[1-9]\d*)(\.\d+)?', number):
        grouped = False
    elif re.fullmatch(r'-?[1-9]\d{0,2}(,\d{3})+(\.\d+)?', number):
        grouped = True
    else:
        return text, None
    # Excel keeps 15 significant digits, so longer ids or account numbers would silently change
    if sum(c.isdigit() for c in number) > 15:
        return text, None

    parsed = float(number.replace(',', ''))
    decimals = len(number.partition('.')[2])
    if percent:
        return parsed / 100, f'0{"." + "0" * decimals if decimals else ""}%'
    if decimals == 0 and not grouped:
        return int(parsed), None
    return parsed, f'{"#,##0" if grouped else "0"}{"." + "0" * decimals if decimals else ""}'


def _build_xlsx_document_bytes(content: str) -> bytes:
    import csv
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    # Every cell without a font of its own uses the workbook's first font, so this sets Google's Arial 10
    wb._fonts[0] = Font(name=XLSX_FONT[0], size=XLSX_FONT[1])
    ws = wb.active
    ws.title = 'Sheet1'
    ws.sheet_format.defaultColWidth = XLSX_MIN_COLUMN_WIDTH
    ws.sheet_format.defaultRowHeight = XLSX_ROW_HEIGHT
    rows = list(csv.reader(io.StringIO(content)))
    widths: dict[int, int] = {}
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row, start=1):
            value, number_format = _xlsx_cell_value(text) if r > 1 else (text, None)
            cell = ws.cell(row=r, column=c, value=value)
            if number_format:
                cell.number_format = number_format
            widths[c] = max(widths.get(c, 0), len(text))

    # A generated table almost always opens with a header, so it's bold and stays in view while scrolling
    if len(rows) > 1:
        for cell in ws[1]:
            cell.font = Font(name=XLSX_FONT[0], size=XLSX_FONT[1], bold=True)
        ws.freeze_panes = 'A2'
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = min(
            max(XLSX_MIN_COLUMN_WIDTH, width * 1.1 + 2), XLSX_MAX_COLUMN_WIDTH
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Google Slides' default 16:9 page (10 x 5.625in) and its "Simple Light" layouts, in EMU as (left, top, width, height)
PPTX_SLIDE_SIZE = (9144000, 5143500)
PPTX_TITLE_BOX = (311700, 445025, 8520600, 572700)
PPTX_BODY_BOX = (311700, 1152475, 8520600, 3416400)
PPTX_COVER_TITLE_BOX = (311700, 744575, 8520600, 2052600)
PPTX_COVER_SUBTITLE_BOX = (311700, 2834125, 8520600, 792600)
PPTX_TITLE_PT = 28
PPTX_BODY_PT = 18
PPTX_COVER_TITLE_PT = 52
PPTX_COVER_SUBTITLE_PT = 28
PPTX_FONT = 'Arial'
# (bullet, hanging indent in EMU) for the first two body levels, as Google Slides draws them
PPTX_BULLETS = (('●', -342900), ('○', -317500))
PPTX_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def _pptx_level_style(shape_element, level: int):
    """A placeholder's own style for one indent level, created when the template leaves it to the master."""
    from lxml import etree
    from pptx.oxml.ns import qn

    lst_style = shape_element.find(qn('p:txBody')).find(qn('a:lstStyle'))
    ppr = lst_style.find(qn(f'a:lvl{level}pPr'))
    if ppr is None:
        ppr = etree.SubElement(lst_style, qn(f'a:lvl{level}pPr'))
    return ppr


def _apply_google_slides_layout(prs) -> None:
    """Turn python-pptx's 4:3 Calibri template into Google Slides' 16:9 Arial one, for new slides too."""
    import re

    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    old_width, old_height = prs.slide_width, prs.slide_height
    prs.slide_width, prs.slide_height = (Emu(v) for v in PPTX_SLIDE_SIZE)

    def place(shape, box):
        shape.left, shape.top, shape.width, shape.height = (Emu(v) for v in box)

    master = prs.slide_master
    cover_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]
    # Everything is first scaled to the new page, then the placeholders we use get Google's exact boxes
    for shapes in [master.placeholders, *(layout.placeholders for layout in prs.slide_layouts)]:
        for shape in shapes:
            if shape._element.spPr.xfrm is not None:
                place(
                    shape,
                    (
                        round(shape.left * prs.slide_width / old_width),
                        round(shape.top * prs.slide_height / old_height),
                        round(shape.width * prs.slide_width / old_width),
                        round(shape.height * prs.slide_height / old_height),
                    ),
                )
    for shapes in (master.placeholders, content_layout.placeholders):
        for shape in shapes:
            if shape.placeholder_format.idx in (0, 1):
                place(shape, PPTX_TITLE_BOX if shape.placeholder_format.idx == 0 else PPTX_BODY_BOX)
    for shape in cover_layout.placeholders:
        if shape.placeholder_format.idx in (0, 1):
            place(shape, PPTX_COVER_TITLE_BOX if shape.placeholder_format.idx == 0 else PPTX_COVER_SUBTITLE_BOX)

    # Default sizes for text that doesn't set its own, every body level at 18pt like Google's
    tx_styles = master._element.find(qn('p:txStyles'))
    for style, size in (('p:titleStyle', PPTX_TITLE_PT), ('p:bodyStyle', PPTX_BODY_PT)):
        for def_rpr in tx_styles.find(qn(style)).iter(qn('a:defRPr')):
            def_rpr.set('sz', str(size * 100))
    # The template centers titles, Google's are left-aligned except on the cover
    tx_styles.find(qn('p:titleStyle')).find(qn('a:lvl1pPr')).set('algn', 'l')
    for shape in cover_layout.placeholders:
        if shape.placeholder_format.idx in (0, 1):
            _pptx_level_style(shape._element, 1).set('algn', 'ctr')

    # Google's bullets: a filled then a hollow circle, text 0.5in in, 115% line spacing and no gap between items
    body_style = tx_styles.find(qn('p:bodyStyle'))
    for level, (char, indent) in enumerate(PPTX_BULLETS, start=1):
        ppr = body_style.find(qn(f'a:lvl{level}pPr'))
        ppr.set('marL', str(457200 * level))
        ppr.set('indent', str(indent))
        ppr.find(qn('a:buChar')).set('char', char)
        for tag in ('a:spcBef', 'a:lnSpc'):
            for old in ppr.findall(qn(tag)):
                ppr.remove(old)
        ppr.insert(0, etree.fromstring(f'<a:spcBef {PPTX_NS}><a:spcPts val="0"/></a:spcBef>'))
        ppr.insert(0, etree.fromstring(f'<a:lnSpc {PPTX_NS}><a:spcPct val="115000"/></a:lnSpc>'))

    # Text sits at the top of its box, except the cover title, which rests on the subtitle
    for shapes in (master.placeholders, content_layout.placeholders, cover_layout.placeholders):
        for shape in shapes:
            if shape.placeholder_format.idx in (0, 1):
                is_cover_title = shapes is cover_layout.placeholders and shape.placeholder_format.idx == 0
                shape._element.find(qn('p:txBody')).find(qn('a:bodyPr')).set(
                    'anchor', 'b' if is_cover_title else 't'
                )

    theme = master.part.part_related_by(RT.THEME)
    theme_xml = theme.blob.decode('utf-8')
    for font in ('majorFont', 'minorFont'):
        theme_xml = re.sub(
            rf'(<a:{font}>\s*<a:latin typeface=")[^"]*', lambda m: m.group(1) + PPTX_FONT, theme_xml, count=1
        )
    theme._blob = theme_xml.encode('utf-8')


def _fill_text_frame(text_frame, lines: list[str], size_pt: int | None = None) -> None:
    from pptx.util import Pt

    text_frame.clear()
    for i, line in enumerate(lines):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        run = p.add_run()
        run.text = line
        if size_pt:
            run.font.size = Pt(size_pt)


def _build_pptx_document_bytes(content: str) -> bytes:
    import io

    from pptx import Presentation

    prs = Presentation()
    _apply_google_slides_layout(prs)
    cover_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]

    chunks = [
        [line.strip() for line in chunk.strip().splitlines() if line.strip()] for chunk in content.split('---')
    ]
    chunks = [chunk for chunk in chunks if chunk]
    # Content like "---" passes the empty check but leaves no slide, and an empty deck isn't a success
    if not chunks:
        raise ValueError('No slides in the content')
    for i, lines in enumerate(chunks):
        title = lines[0].lstrip('#').strip()
        rest = [line.lstrip('-*').strip() for line in lines[1:]]

        # An opening slide with just a title and at most a subtitle reads as a cover, like Google's "Title slide"
        if i == 0 and len(rest) <= 1:
            slide = prs.slides.add_slide(cover_layout)
            _fill_text_frame(slide.shapes.title.text_frame, [title], PPTX_COVER_TITLE_PT)
            subtitle = slide.placeholders[1]
            if rest:
                _fill_text_frame(subtitle.text_frame, rest, PPTX_COVER_SUBTITLE_PT)
            else:
                subtitle._element.getparent().remove(subtitle._element)
            continue

        slide = prs.slides.add_slide(content_layout)
        _fill_text_frame(slide.shapes.title.text_frame, [title])
        if rest:
            _fill_text_frame(slide.placeholders[1].text_frame, rest)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_document_bytes(format: str, name: str, content: str) -> bytes:
    if format == 'pdf':
        return _build_pdf_document_bytes(name, content)
    elif format == 'docx':
        return _build_docx_document_bytes(content)
    elif format == 'xlsx':
        return _build_xlsx_document_bytes(content)
    return _build_pptx_document_bytes(content)



def main() -> int:
    global _reference_doc, _skip_reference_doc
    spec = json.loads(sys.stdin.buffer.read())
    for key in ('fonts_dir', 'tmp_dir'):
        if spec.get(key):
            _paths[key] = spec[key]
    # The parent built the restyled reference.docx once, so each child doesn't build it again
    if 'reference_doc' in spec:
        if spec['reference_doc']:
            _reference_doc = Path(spec['reference_doc'])
        else:
            _skip_reference_doc = True
    out = sys.stdout.buffer
    # Libraries may print to stdout, and the file must be the only thing there
    sys.stdout = sys.stderr
    try:
        file_bytes = build_document_bytes(spec['format'], spec['name'], spec['content'])
    except Exception:
        log.exception(f'Building the {spec.get("format")} document failed')
        return 1
    if len(file_bytes) > spec.get('max_bytes', float('inf')):
        return EXIT_TOO_LARGE
    out.write(file_bytes)
    out.flush()
    return 0


if __name__ == '__main__':
    sys.exit(main())
